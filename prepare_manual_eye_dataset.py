#!/usr/bin/env python3
"""Validate and prepare manual eye crops while retaining unmodified originals.

Author: Jia Rong

The clinical class and the visibility/eligibility of a particular crop are
separate fields. This utility never infers a clinical diagnosis from a photo.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import random
import re
import zipfile
from collections import Counter, defaultdict
from contextlib import ExitStack
from pathlib import Path

from PIL import Image, ImageDraw, ImageOps, ImageStat

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path, rows, fields=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(dict.fromkeys(k for row in rows for k in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def hashes(path):
    md5 = hashlib.md5()
    sha = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            md5.update(chunk)
            sha.update(chunk)
    return md5.hexdigest(), sha.hexdigest()


def image_info(path):
    with Image.open(path) as image:
        image.load()
        width, height = image.size
        rotation = int(image.getexif().get(274, 1))
        small = image.convert("RGB")
        small.thumbnail((256, 256))
        gray = small.convert("L")
        st = ImageStat.Stat(gray)
        quality = []
        if min(width, height) < 200:
            quality.append("small_native_image")
        if width / max(height, 1) < 0.55 or width / max(height, 1) > 3.0:
            quality.append("extreme_aspect_ratio")
        if st.mean[0] < 30:
            quality.append("dark_image")
        if st.stddev[0] < 8:
            quality.append("low_contrast")
        return dict(width=width, height=height, exif_orientation=rotation,
                    quality_flags=";".join(quality))


def safe_relative(base, relative):
    parts = Path(relative.replace("\\", "/"))
    if parts.is_absolute() or ".." in parts.parts:
        raise ValueError(f"Unsafe relative path: {relative}")
    resolved = (base / parts).resolve()
    if not resolved.is_relative_to(base.resolve()):
        raise ValueError(f"Path leaves dataset: {relative}")
    return resolved


def case_id(relative):
    # Prefer the nearest case directory, not a collection range such as
    # 1500-1985, a file name, or a dated follow-up subdirectory.
    for part in reversed(Path(relative.replace("\\", "/")).parts[:-1]):
        if re.fullmatch(r"\d+\s*[-_－—]\s*\d+", part):
            continue
        match = re.match(r"^(\d{2,6})\s*[-_－—]", part)
        if match:
            return match.group(1)
    return ""


def normalized_prior_patients(rows):
    by_case = defaultdict(list)
    for row in rows:
        cid = case_id(row["source_relative_path"])
        if cid:
            by_case[cid].append(row)
    aliases = {}
    for cid, items in by_case.items():
        labels = {int(r["label"]) for r in items}
        if len(labels) != 1:
            raise ValueError(f"Clinical case has conflicting class labels: {cid}")
        groups = {r["patient_group"] for r in items}
        numbered = sorted(g for g in groups if g.startswith(f"raw_patient_num_{cid}_"))
        canonical = numbered[0] if numbered else sorted(groups)[0]
        for group in groups:
            previous = aliases.get(group)
            if previous is not None and previous != canonical:
                raise ValueError(f"Source group maps to multiple clinical cases: {group}")
            aliases[group] = canonical
    normalized = [dict(row, source_patient_group=row["patient_group"],
                       patient_group=aliases.get(row["patient_group"], row["patient_group"]))
                  for row in rows]
    return normalized, {group: target for group, target in aliases.items() if group != target}


def panel_preview(rows, base, destination, prefix="batch07", per_page=60):
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    cell_w, cell_h, columns = 230, 175, 10
    index = []
    for start in range(0, len(rows), per_page):
        page_rows = rows[start:start + per_page]
        height = ((len(page_rows) + columns - 1) // columns) * cell_h
        canvas = Image.new("RGB", (columns * cell_w, height), "white")
        draw = ImageDraw.Draw(canvas)
        for local_index, row in enumerate(page_rows):
            slot = start + local_index + 1
            x, y = (local_index % columns) * cell_w, (local_index // columns) * cell_h
            try:
                with Image.open(safe_relative(base, row["source_crop_path"])) as original:
                    thumb = ImageOps.exif_transpose(original).convert("RGB")
                    thumb.thumbnail((cell_w - 8, cell_h - 28))
                    canvas.paste(thumb, (x + (cell_w - thumb.width) // 2, y + 4))
            except Exception as exc:
                draw.text((x + 5, y + 20), f"DECODE ERROR {type(exc).__name__}", fill="red")
            tag = f"{slot:04d} case {row.get('case_id', '')} {'M' if row['label'] == 1 else 'B'}"
            draw.text((x + 5, y + cell_h - 21), tag, fill="black")
            index.append({"preview_index": slot, "page": start // per_page + 1,
                          "source_crop_path": row["source_crop_path"],
                          "source_original_path": row["source_original_path"]})
        canvas.save(destination / f"{prefix}_{start // per_page + 1:02d}.jpg", quality=92)
    write_csv(destination / "preview_index.csv", index)


def preview_catalog(args):
    dataset = Path(args.dataset_root).resolve()
    rows = read_csv(dataset / "metadata" / "crops.csv")
    if args.family == "main":
        rows = [r for r in rows if r["cohort"] == "prior_cohort" and r["eligibility"] == "manual_crop_available"]
    elif args.family == "uncertain":
        rows = [r for r in rows if r["eligibility"] != "manual_crop_available"]
    elif args.family == "missing":
        rows = read_csv(dataset / "metadata" / "originals_without_current_crop.csv")
    output = Path(args.output).resolve()
    if not output.is_relative_to(dataset.parent.parent / "tmp"):
        raise ValueError("Preview output must be in this project's tmp subtree")
    output.mkdir(parents=True, exist_ok=True)
    index = []
    for start in range(0, len(rows), 60):
        page = rows[start:start+60]
        canvas = Image.new("RGB", (2300, 175 * ((len(page)+9)//10)), "white")
        draw = ImageDraw.Draw(canvas)
        for j, row in enumerate(page):
            relative = row.get("crop_path", row["original_path"])
            with Image.open(safe_relative(dataset, relative)) as image:
                thumb = ImageOps.exif_transpose(image).convert("RGB")
                thumb.thumbnail((222, 147))
                x, y = (j % 10) * 230, (j // 10) * 175
                canvas.paste(thumb, (x + (230-thumb.width)//2, y+4))
            n = start+j+1
            draw.text((x+4, y+154), f"{n:04d} {Path(relative).stem[-22:]} {row['label']}", fill="black")
            index.append(dict(preview_index=n, family=args.family, path=relative,
                              original_path=row["original_path"], original_id=row["original_id"],
                              patient_group=row["patient_group"], label=row["label"],
                              source_original_path=row.get("source_original_path", "")))
        canvas.save(output / f"{args.family}_{start//60+1:02d}.jpg", quality=92)
    write_csv(output / f"{args.family}_index.csv", index)
    print(json.dumps({"family": args.family, "images": len(rows), "pages": (len(rows)+59)//60}))


def audit_new(args):
    from openpyxl import load_workbook

    base = Path(args.workspace).resolve()
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    old, patient_aliases = normalized_prior_patients(read_csv(args.old_manifest))
    old_hash = defaultdict(list)
    old_cases = defaultdict(list)
    for row in old:
        old_hash[row["md5"]].append(row)
        cid = case_id(row.get("source_relative_path", ""))
        if cid:
            old_cases[cid].append(row)
    originals, crops, reviews, issues = [], [], [], []
    original_by_path = {}
    supplied_files = set()
    for category, label in [("恶性", 1), ("良性", 0)]:
        raw_root = base / category
        crop_root = base / f"{category}cropped"
        if not raw_root.exists():
            raw_root = base / "eye_dataset" / "originals" / "batch07" / category
            crop_root = base / "eye_dataset" / "cropped" / "batch07" / f"{category}cropped"
        for path in sorted(raw_root.rglob("*")):
            if path.suffix.lower() not in IMAGE_EXTS or not path.is_file():
                continue
            relative = path.relative_to(base).as_posix()
            source_relative = f"{category}/" + path.relative_to(raw_root).as_posix()
            md5, sha = hashes(path)
            row = dict(source_original_path=relative, source_lookup_path=source_relative,
                       case_id=case_id(source_relative), label=label,
                       class_name="malignant" if label else "benign", md5=md5, sha256=sha,
                       bytes=path.stat().st_size, label_source="supplied_clinical_class_folder",
                       cohort="batch07", intended_role="temporal_validation",
                       collection_date="", date_evidence="not_supplied")
            try:
                row.update(image_info(path))
            except Exception as exc:
                row["decode_error"] = str(exc)
                issues.append(dict(kind="original_decode_error", path=relative, detail=str(exc)))
            matched = old_hash.get(md5, [])
            row["prior_image_ids"] = ";".join(r["image_id"] for r in matched)
            row["prior_patient_groups"] = ";".join(sorted({r["patient_group"] for r in matched}))
            if any(int(r["label"]) != label for r in matched):
                issues.append(dict(kind="label_conflict_with_prior_source", path=relative, detail=row["prior_image_ids"]))
            original_by_path[source_relative] = row
            originals.append(row)
        mapping = read_csv(crop_root / "crop_mapping.csv")
        review_paths = sorted(crop_root.rglob("review_needed.csv"))
        for path in review_paths:
            for item in read_csv(path):
                reviews.append(dict(item, category=category,
                                    source_review_path=path.relative_to(base).as_posix()))
        seen_mapping = set()
        for item in mapping:
            source_crop = item["cropped_filename"].replace("\\", "/")
            crop_relative = source_crop.split("/", 1)[1]
            path = (crop_root / crop_relative).resolve()
            if not path.is_relative_to(crop_root.resolve()):
                raise ValueError(f"Crop mapping leaves supplied folder: {source_crop}")
            relative = path.relative_to(base).as_posix()
            if relative in seen_mapping:
                issues.append(dict(kind="duplicate_mapping_row", path=relative, detail=""))
            seen_mapping.add(relative)
            original = original_by_path.get(item["original_filename"].replace("\\", "/"))
            if original is None:
                issues.append(dict(kind="missing_original_mapping", path=relative, detail=item["original_filename"]))
                continue
            if not path.is_file():
                issues.append(dict(kind="missing_crop_file", path=relative, detail=""))
                continue
            md5, sha = hashes(path)
            row = dict(source_crop_path=relative, source_original_path=original["source_original_path"],
                       source_lookup_original=original["source_lookup_path"], case_id=original["case_id"],
                       label=label, class_name=original["class_name"], original_sha256=original["sha256"],
                       original_md5=original["md5"], crop_sha256=sha, crop_md5=md5,
                       bytes=path.stat().st_size, cohort="batch07", intended_role="temporal_validation",
                       lesion_visibility="confirmed_by_supplied_manual_mapping",
                       label_source=original["label_source"], crop_type=item["crop_type"],
                       mapping_notes=item.get("notes", ""), eligibility="manual_crop_available")
            for field in ["crop_x", "crop_y", "crop_width", "crop_height", "output_width", "output_height", "aspect_ratio"]:
                row[field] = item.get(field, "")
            try:
                row.update(image_info(path))
                if (row["width"], row["height"]) != (int(item["output_width"]), int(item["output_height"])):
                    issues.append(dict(kind="output_dimension_mismatch", path=relative,
                                       detail=f"file={row['width']}x{row['height']} csv={item['output_width']}x{item['output_height']}"))
                x, y, w, h = [int(item[f]) for f in ["crop_x", "crop_y", "crop_width", "crop_height"]]
                ow, oh = original["width"], original["height"]
                if min(x, y) < 0 or min(w, h) <= 0 or not ((x+w <= ow and y+h <= oh) or (x+w <= oh and y+h <= ow)):
                    issues.append(dict(kind="crop_coordinates_out_of_bounds", path=relative,
                                       detail=f"box={x},{y},{w},{h} original={ow},{oh}"))
                if (w, h) != (row["width"], row["height"]):
                    row["resize_record"] = "mapping_indicates_resampling"
                else:
                    row["resize_record"] = "native_crop_dimensions_preserved"
            except Exception as exc:
                row["decode_error"] = str(exc)
                issues.append(dict(kind="crop_decode_error", path=relative, detail=str(exc)))
            supplied_files.add(relative)
            crops.append(row)
        for path in crop_root.rglob("*"):
            if path.is_file() and path.suffix.lower() in IMAGE_EXTS and path.relative_to(base).as_posix() not in seen_mapping:
                issues.append(dict(kind="unmapped_crop_file", path=path.relative_to(base).as_posix(), detail=""))
    workbook_path = base / "第七批裁剪总表.xlsx"
    if not workbook_path.is_file():
        workbook_path = output / workbook_path.name
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    book_counts = {sheet.title: sum(1 for _ in sheet.iter_rows(values_only=True)) - 1 for sheet in book}
    sheet = book["全部裁图清单"]
    values = sheet.iter_rows(values_only=True)
    headers = next(values)
    xlsx_crops = {dict(zip(headers, v))["裁图相对路径"].replace("\\", "/") for v in values}
    csv_crops = {r["source_crop_path"].split("cropped/batch07/", 1)[-1] for r in crops}
    for relative in sorted(xlsx_crops.symmetric_difference(csv_crops)):
        issues.append(dict(kind="xlsx_csv_crop_disagreement", path=relative, detail=""))
    book.close()
    new_cases = {r["case_id"] for r in originals}
    overlapping_groups = {r["patient_group"] for cid in new_cases for r in old_cases.get(cid, [])}
    overlapping_groups.update(g for r in originals for g in r["prior_patient_groups"].split(";") if g)
    removed = [dict(r, removal_reason="patient_present_in_batch07_temporal_validation")
               for r in old if r["patient_group"] in overlapping_groups]
    for cid in sorted(new_cases):
        labels = {r["label"] for r in originals if r["case_id"] == cid}
        if len(labels) != 1:
            issues.append(dict(kind="mixed_new_patient_labels", path=cid, detail=str(sorted(labels))))
    group_hash = defaultdict(list)
    crop_hash = defaultdict(list)
    for row in originals:
        group_hash[row["sha256"]].append(row)
    for row in crops:
        crop_hash[row["crop_sha256"]].append(row)
    for digest, items in group_hash.items():
        if len({(r["case_id"], r["label"]) for r in items}) > 1:
            issues.append(dict(kind="identical_original_cross_patient", path=digest, detail=";".join(r["source_original_path"] for r in items)))
    used_originals = {r["source_original_path"] for r in crops}
    not_cropped = [r for r in originals if r["source_original_path"] not in used_originals]
    summary = dict(originals=len(originals), crops=len(crops), patient_cases=len(new_cases),
                   original_class_counts=dict(Counter(r["class_name"] for r in originals)),
                   crop_class_counts=dict(Counter(r["class_name"] for r in crops)),
                   original_without_crop=len(not_cropped), doctor_review_records=len(reviews),
                   doctor_review_status=dict(Counter(r.get("处理状态", "") for r in reviews)),
                   workbook_data_rows=book_counts,
                   prior_exact_original_matches=sum(bool(r["prior_image_ids"]) for r in originals),
                   prior_overlapping_patient_groups=len(overlapping_groups),
                   prior_records_to_remove_from_development=len(removed),
                   repeated_crop_content_groups=sum(len(v) > 1 for v in crop_hash.values()),
                   technical_issues=len(issues), issue_counts=dict(Counter(r["kind"] for r in issues)),
                   quality_flagged_crops=sum(bool(r.get("quality_flags")) for r in crops),
                   cohort_source="same_hospital", intended_validation="temporal",
                   chronology_dates_verified=False,
                   clinical_label_note="Labels are copied from supplied clinical class folders, not inferred from image appearance.")
    write_csv(output / "batch07_originals.csv", originals)
    write_csv(output / "batch07_crops.csv", crops)
    write_csv(output / "batch07_doctor_review.csv", reviews)
    write_csv(output / "batch07_originals_without_crop.csv", not_cropped)
    write_csv(output / "prior_development_exclusions.csv", removed)
    write_csv(output / "technical_issues.csv", issues, ["kind", "path", "detail"])
    write_json(output / "batch07_validation.json", summary)
    if args.preview_dir:
        panel_preview(crops, base, args.preview_dir)
    print(json.dumps(summary, ensure_ascii=False))
    return 1 if issues else 0


def hidden_zip_entry(name):
    return any(p == "__MACOSX" or p.startswith("._") for p in Path(name).parts)


def zip_rows(archive, suffix):
    candidates = [n for n in archive.namelist() if n.endswith(suffix) and not hidden_zip_entry(n)]
    if len(candidates) != 1:
        raise ValueError(f"Expected one {suffix}: found {candidates}")
    return list(csv.DictReader(io.StringIO(archive.read(candidates[0]).decode("utf-8-sig"))))


def select_manual_packages(directory):
    """Resolve one returned archive per part, preferring the supplied ' 2' files."""
    grouped = defaultdict(list)
    for path in Path(directory).glob("eye_crop_manual_review_part_*_of_05*.zip"):
        match = re.fullmatch(r"eye_crop_manual_review_part_(\d{2})_of_05(?: 2)?\.zip", path.name)
        if match:
            grouped[int(match.group(1))].append(path)
    if set(grouped) != set(range(1, 6)):
        raise ValueError("Exactly the five manual-crop archive parts 01 through 05 are required")
    selected = []
    for part in range(1, 6):
        returned = [p for p in grouped[part] if p.name.endswith(" 2.zip")]
        candidates = returned or grouped[part]
        if len(candidates) != 1:
            raise ValueError(f"Ambiguous returned manual-crop archive part {part:02d}")
        selected.append(candidates[0])
    return selected


def check_crop_creation_policy(dataset):
    design_path = Path(dataset) / "metadata" / "cohort_design.json"
    if not design_path.is_file():
        raise ValueError("A current cohort design is required before creating additional crops")
    design = json.loads(design_path.read_text(encoding="utf-8"))
    if design.get("crop_source_policy") != "supplement_missing_originals":
        raise ValueError("Current crop policy reuses existing manual crops; new crop creation is disabled")


def link_file(source, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if hashes(source)[1] != hashes(destination)[1]:
            raise ValueError(f"Existing destination differs: {destination}")
        return
    os.link(source, destination)


def consolidate_remote(args):
    root = Path(args.project_root).resolve()
    dataset = Path(args.dataset_root).resolve()
    if not dataset.is_relative_to(root / "data_stream"):
        raise ValueError("Dataset must be inside this project's data_stream directory")
    metadata = dataset / "metadata"
    new_raw = read_csv(metadata / "batch07_originals.csv")
    new_crops = read_csv(metadata / "batch07_crops.csv")
    excluded = read_csv(metadata / "prior_development_exclusions.csv")
    old, patient_aliases = normalized_prior_patients(read_csv(args.old_manifest))
    removed_groups = {r["patient_group"] for r in excluded}
    old_hash = defaultdict(list)
    old_case_groups = defaultdict(set)
    for row in old:
        old_hash[row["md5"]].append(row)
        cid = case_id(row.get("source_relative_path", ""))
        if cid:
            old_case_groups[cid].add(row["patient_group"])
    originals, crop_rows, issues, not_selected = [], [], [], []
    digest_catalog = {}
    new_by_md5 = {}
    for row in new_raw:
        label = int(row["label"])
        path = dataset / "originals" / "batch07" / row["source_lookup_path"]
        if not path.is_file():
            raise FileNotFoundError(path)
        md5, sha = hashes(path)
        if sha != row["sha256"] or md5 != row["md5"]:
            raise ValueError(f"Uploaded original differs: {path}")
        matched_groups = sorted({r["patient_group"] for r in old_hash.get(md5, [])}
                                | old_case_groups.get(row["case_id"], set()))
        if len(matched_groups) > 1:
            raise ValueError(f"Ambiguous prior patient identity for {path}")
        group = matched_groups[0] if matched_groups else f"batch07_patient_{row['case_id']}"
        info = image_info(path)
        record = dict(original_id=f"img_{sha[:20]}", original_path=path.relative_to(dataset).as_posix(),
                      patient_group=group, case_id=row["case_id"], label=label,
                      class_name=row["class_name"], cohort="batch07", intended_role="temporal_validation",
                      md5=md5, sha256=sha, bytes=path.stat().st_size,
                      source_image_ids=row["prior_image_ids"], source_original_path=row["source_lookup_path"],
                      label_source=row["label_source"], collection_date="", **info)
        if sha in digest_catalog:
            if (digest_catalog[sha]["patient_group"], digest_catalog[sha]["label"]) != (group, label):
                raise ValueError(f"Exact duplicate with conflicting patient/label: {path}")
            not_selected.append(dict(source_path=str(path), reason="identical_supplied_original", canonical_path=digest_catalog[sha]["original_path"]))
        else:
            originals.append(record)
            digest_catalog[sha] = record
        new_by_md5[md5] = digest_catalog[sha]
    old_originals = {}
    for row in old:
        if row["md5"] in new_by_md5:
            original = new_by_md5[row["md5"]]
            if int(row["label"]) != original["label"]:
                raise ValueError(f"Clinical label conflict for {row['image_id']}")
            old_originals[row["image_id"]] = original
            continue
        source = root / "data_stream" / "raw_tumor_dataset" / row["canonical_relative_path"]
        if not source.is_file():
            raise FileNotFoundError(source)
        md5, sha = hashes(source)
        if md5 != row["md5"]:
            raise ValueError(f"Prior source MD5 differs: {row['image_id']}")
        group = row["patient_group"]
        role = "temporal_validation" if group in removed_groups else "development"
        if sha in digest_catalog:
            original = digest_catalog[sha]
            if (original["patient_group"], original["label"]) != (group, int(row["label"])):
                issues.append(dict(kind="prior_duplicate_patient_conflict", path=row["image_id"], detail=original["original_id"]))
        else:
            destination = dataset / "originals" / "prior_cohort" / f"{row['image_id']}{source.suffix.lower()}"
            link_file(source, destination)
            original = dict(original_id=f"img_{sha[:20]}", original_path=destination.relative_to(dataset).as_posix(),
                            patient_group=group, case_id=case_id(row.get("source_relative_path", "")),
                            label=int(row["label"]), class_name=row["class_name"], cohort="prior_cohort",
                            intended_role=role, md5=md5, sha256=sha, bytes=source.stat().st_size,
                            source_image_ids=row["image_id"], source_original_path=row.get("source_relative_path", ""),
                            label_source="prior_clinical_manifest", collection_date="", **image_info(source))
            originals.append(original)
            digest_catalog[sha] = original
        old_originals[row["image_id"]] = original
    # Retain the verified catalog even if a later archive row needs repair.
    write_csv(metadata / "originals.csv", originals)
    crop_hashes = {}
    doctor_reviews = read_csv(metadata / "batch07_doctor_review.csv")
    for row in new_crops:
        relative = row["source_crop_path"].split("cropped/batch07/", 1)[-1]
        path = dataset / "cropped" / "batch07" / relative
        md5, sha = hashes(path)
        if sha != row["crop_sha256"]:
            raise ValueError(f"Uploaded crop differs: {path}")
        original = new_by_md5[row["original_md5"]]
        crop_doctor_reviews = [r for r in doctor_reviews
                              if r["original_filename"].replace("\\", "/") == row["source_lookup_original"]]
        record = dict(crop_id=f"crop_{sha[:20]}", crop_path=path.relative_to(dataset).as_posix(),
                      original_id=original["original_id"], original_path=original["original_path"],
                      patient_group=original["patient_group"], label=original["label"],
                      class_name=original["class_name"], cohort="batch07", intended_role="temporal_validation",
                      lesion_visibility=row["lesion_visibility"], eligibility="manual_crop_available",
                      clinical_review="explicit_supplied_doctor_review" if crop_doctor_reviews else "supplied_manual_mapping",
                      doctor_review_status=";".join(r.get("处理状态", "") for r in crop_doctor_reviews),
                      doctor_review_comment=";".join(r.get("审核意见", "") for r in crop_doctor_reviews),
                      crop_type=row["crop_type"],
                      sha256=sha, md5=md5, label_source=row["label_source"],
                      source_original_path=row["source_lookup_original"], source_crop_path=relative,
                      source_package="batch07_supplied_folders", notes=row["mapping_notes"],
                      **image_info(path))
        for field in ["crop_x", "crop_y", "crop_width", "crop_height", "output_width", "output_height", "aspect_ratio"]:
            record[field] = row[field]
        if sha in crop_hashes:
            previous = crop_hashes[sha]
            if (previous["original_id"], previous["label"]) != (record["original_id"], record["label"]):
                issues.append(dict(kind="crop_duplicate_source_conflict", path=record["crop_path"], detail=previous["crop_path"]))
            not_selected.append(dict(source_path=str(path), reason="identical_supplied_crop", canonical_path=previous["crop_path"]))
        else:
            crop_rows.append(record)
            crop_hashes[sha] = record
    packages = select_manual_packages(args.manual_packages)
    raw_covered = set()
    superseded = []
    uncertain_reasons = {}
    for package in packages:
        with zipfile.ZipFile(package) as archive:
            members = [n for n in archive.namelist() if not hidden_zip_entry(n)]
            for item in zip_rows(archive, "/cropped/review_needed.csv"):
                uncertain_reasons[item["original_filename"]] = item.get("reason", "")
            original_manifest = zip_rows(archive, "/manual_crop_manifest.csv")
            raw_covered.update(r["image_id"] for r in original_manifest)
            main_mapping = zip_rows(archive, "/cropped/crop_mapping.csv")
            main_filenames = {Path(r["cropped_filename"]).name for r in main_mapping}
            for suffix, status in [("/cropped/crop_mapping.csv", "manual_crop_available"),
                                   ("/cropped/uncertain_lesion/crop_mapping.csv", "needs_clinical_lesion_review")]:
                for item in zip_rows(archive, suffix):
                    stem = Path(item["original_filename"]).stem
                    original = old_originals.get(stem)
                    if original is None:
                        raise ValueError(f"Unknown prior original ID: {stem}")
                    if original["cohort"] == "batch07":
                        superseded.append(dict(source_package=package.name, source_crop=item["cropped_filename"],
                                               original_id=original["original_id"], reason="superseded_by_batch07_doctor_reviewed_mapping"))
                        continue
                    filename = Path(item["cropped_filename"]).name
                    if status == "needs_clinical_lesion_review" and filename in main_filenames:
                        superseded.append(dict(source_package=package.name, source_crop=item["cropped_filename"],
                                               original_id=original["original_id"], reason="older_uncertain_crop_replaced_by_current_main_manual_mapping"))
                        continue
                    wanted_suffix = ("/cropped/uncertain_lesion/" if status == "needs_clinical_lesion_review" else "/cropped/") + filename
                    found = [n for n in members if n.endswith(wanted_suffix)]
                    if len(found) != 1:
                        raise ValueError(f"Ambiguous or absent crop in {package.name}: {filename}")
                    data = archive.read(found[0])
                    md5, sha = hashlib.md5(data).hexdigest(), hashlib.sha256(data).hexdigest()
                    if sha in crop_hashes:
                        previous = crop_hashes[sha]
                        if previous["original_id"] != original["original_id"]:
                            issues.append(dict(kind="crop_duplicate_source_conflict", path=filename, detail=previous["crop_path"]))
                        not_selected.append(dict(source_path=f"{package.name}:{found[0]}", reason="identical_crop_content", canonical_path=previous["crop_path"]))
                        continue
                    destination = dataset / "cropped" / "prior_cohort" / filename
                    destination.parent.mkdir(parents=True, exist_ok=True)
                    if destination.exists():
                        if hashes(destination)[1] != sha:
                            raise ValueError(f"Existing prior crop differs: {destination}")
                    else:
                        destination.write_bytes(data)
                    info = image_info(destination)
                    if (info["width"], info["height"]) != (int(item["output_width"]), int(item["output_height"])):
                        issues.append(dict(kind="prior_crop_dimension_mismatch", path=filename, detail=str(info)))
                    record = dict(crop_id=f"crop_{sha[:20]}", crop_path=destination.relative_to(dataset).as_posix(),
                                  original_id=original["original_id"], original_path=original["original_path"],
                                  patient_group=original["patient_group"], label=original["label"],
                                  class_name=original["class_name"], cohort="prior_cohort",
                                  intended_role=original["intended_role"], crop_type=item["crop_type"],
                                  lesion_visibility="uncertain" if status == "needs_clinical_lesion_review" else "confirmed_by_supplied_manual_mapping",
                                  eligibility=status, clinical_review="manual_cropper_mapping_only",
                                  sha256=sha, md5=md5, label_source=original["label_source"],
                                  source_original_path=item["original_filename"], source_crop_path=item["cropped_filename"],
                                  source_package=package.name, notes=item.get("notes", ""),
                                  review_reason=uncertain_reasons.get(item["original_filename"], ""), **info)
                    for field in ["crop_x", "crop_y", "crop_width", "crop_height", "output_width", "output_height", "aspect_ratio"]:
                        record[field] = item.get(field, "")
                    crop_rows.append(record)
                    crop_hashes[sha] = record
        print(json.dumps({"processed_package": package.name, "current_crops": len(crop_rows)}, ensure_ascii=False), flush=True)
    crop_count = Counter(r["original_id"] for r in crop_rows)
    missing = [dict(r, reason="no_current_manual_crop") for r in originals if not crop_count[r["original_id"]]]
    used_development = [r for r in crop_rows if r["intended_role"] == "development" and r["eligibility"] == "manual_crop_available"]
    used_temporal = [r for r in crop_rows if r["intended_role"] == "temporal_validation" and r["eligibility"] == "manual_crop_available"]
    pending = [r for r in crop_rows if r["eligibility"] != "manual_crop_available"]
    crossing_groups = {r["patient_group"] for r in used_development} & {r["patient_group"] for r in used_temporal}
    crossing_originals = {r["original_id"] for r in used_development} & {r["original_id"] for r in used_temporal}
    crossing_crop_hash = {r["sha256"] for r in used_development} & {r["sha256"] for r in used_temporal}
    if crossing_groups or crossing_originals or crossing_crop_hash:
        raise ValueError("Development/temporal identity crossing detected")
    summary = dict(original_records=len(originals), crop_records=len(crop_rows),
                   development_crops=len(used_development), development_patients=len({r["patient_group"] for r in used_development}),
                   temporal_validation_crops=len(used_temporal), temporal_validation_patients=len({r["patient_group"] for r in used_temporal}),
                   pending_clinical_lesion_review=len(pending), originals_without_current_crop=len(missing),
                   superseded_prior_crops=len(superseded), identical_content_not_selected=len(not_selected),
                   prior_originals_covered_by_returned_manual_packages=len(raw_covered),
                   development_temporal_patient_crossings=len(crossing_groups),
                   development_temporal_original_crossings=len(crossing_originals),
                   development_temporal_exact_crop_crossings=len(crossing_crop_hash),
                   clinical_class_counts=dict(Counter(r["class_name"] for r in crop_rows)),
                   technical_issues=len(issues), issue_counts=dict(Counter(r["kind"] for r in issues)),
                   same_hospital=True, validation_design="temporal",
                   chronological_cutoff_verified=False,
                   note="No training launched. Missing clinical confirmation is not replaced by visual diagnosis. Original pixel data preserved.")
    write_csv(metadata / "originals.csv", originals)
    write_csv(metadata / "crops.csv", crop_rows)
    write_csv(metadata / "development_crops.csv", used_development)
    write_csv(metadata / "temporal_validation_crops.csv", used_temporal)
    write_csv(metadata / "clinical_review_pending.csv", pending)
    write_csv(metadata / "originals_without_current_crop.csv", missing)
    write_csv(metadata / "superseded_prior_crops.csv", superseded)
    write_csv(metadata / "identical_content_not_selected.csv", not_selected, ["source_path", "reason", "canonical_path"])
    write_csv(metadata / "consolidation_issues.csv", issues, ["kind", "path", "detail"])
    write_json(metadata / "dataset_validation.json", summary)
    print(json.dumps(summary, ensure_ascii=False), flush=True)
    return 1 if issues else 0


SPLIT_NAMES = ("train", "validation", "test")


def grouped_partition(rows, seed=42):
    """Assign patient groups, never photos, to an 80/10/10 class-stratified split."""
    groups = defaultdict(list)
    for row in rows:
        group = row.get("patient_group", "")
        if not group:
            raise ValueError("Missing patient identity")
        if int(row["label"]) not in (0, 1):
            raise ValueError("Clinical label must be 0 or 1")
        groups[group].append(row)
    by_class = defaultdict(list)
    for group, items in groups.items():
        labels = {int(r["label"]) for r in items}
        if len(labels) != 1:
            raise ValueError(f"Conflicting clinical labels within patient group: {group}")
        by_class[next(iter(labels))].append(group)
    allocation = {}
    for label, patients in sorted(by_class.items()):
        if len(patients) < 3:
            raise ValueError(f"Class {label} has fewer than three patient groups")
        patients = sorted(patients)
        rng = random.Random(f"{seed}:class:{label}")
        rng.shuffle(patients)
        # Largest-remainder integer allocation. Ratios refer to patient counts;
        # the photograph counts are reported separately and need not be exact.
        fractions = (0.8, 0.1, 0.1)
        exact = [len(patients) * f for f in fractions]
        counts = [int(n) for n in exact]
        remainder_order = sorted(range(3), key=lambda i: (-(exact[i] - counts[i]), i))
        for i in remainder_order[:len(patients) - sum(counts)]:
            counts[i] += 1
        if min(counts) == 0:
            raise ValueError(f"Class {label} is too small for the requested split")
        offset = 0
        for split, count in zip(SPLIT_NAMES, counts):
            for group in patients[offset:offset + count]:
                allocation[group] = split
            offset += count
    return allocation


def assert_partition_integrity(rows):
    """Check identities and identical originals/crops across all four partitions."""
    seen = {field: defaultdict(set) for field in
            ("patient_group", "original_md5", "crop_sha256")}
    errors = []
    image_ids = set()
    clinical_labels = defaultdict(set)
    for row in rows:
        image_id = row["image_id"]
        if image_id in image_ids:
            errors.append(f"Repeated crop record: {image_id}")
        image_ids.add(image_id)
        clinical_labels[row["patient_group"]].add(int(row["label"]))
        split = row["split"]
        if split not in (*SPLIT_NAMES, "prospective_test"):
            errors.append(f"Invalid partition: {split}")
        if row["cohort"] == "batch07" and split != "prospective_test":
            errors.append(f"Batch seven used in development: {image_id}")
        if row["cohort"] != "batch07" and split == "prospective_test":
            errors.append(f"Earlier image added to batch-seven test: {image_id}")
        for field, values in seen.items():
            value = row.get(field, "")
            if not value:
                errors.append(f"Missing {field}: {image_id}")
            else:
                values[value].add(split)
    crossings = {field: sum(len(splits) > 1 for splits in values.values())
                 for field, values in seen.items()}
    errors.extend(f"{field} has {count} partition crossings"
                  for field, count in crossings.items() if count)
    errors.extend(f"Conflicting labels for patient group: {group}"
                  for group, labels in clinical_labels.items() if len(labels) != 1)
    if errors:
        raise ValueError("; ".join(errors[:20]))
    return crossings


def prior_batch(row):
    first = row["source_relative_path"].replace("\\", "/").split("/")[0]
    for batch, token in [(1, "第一批"), (2, "第二批"), (3, "第三批"),
                         (4, "第四批"), (5, "第五批"), (6, "第6批")]:
        if token in first:
            return batch
    raise ValueError(f"Unrecognized prior collection batch: {first}")


def package_crop_records(packages, old_by_id, overlap_groups, quality_decisions):
    records, excluded, covered = [], [], set()
    for package in packages:
        with zipfile.ZipFile(package) as archive:
            members = {n for n in archive.namelist() if not hidden_zip_entry(n)}
            original_manifest = zip_rows(archive, "/manual_crop_manifest.csv")
            for row in original_manifest:
                image_id = row["image_id"]
                if image_id not in old_by_id:
                    raise ValueError(f"Unknown original: {image_id}")
                original = old_by_id[image_id]
                if (row["patient_group"], int(row["label"])) != (original.get("source_patient_group", original["patient_group"]), int(original["label"])):
                    raise ValueError(f"Package patient/label mismatch: {image_id}")
                covered.add(image_id)
            main = zip_rows(archive, "/cropped/crop_mapping.csv")
            main_names = {Path(r["cropped_filename"]).name for r in main}
            uncertain = zip_rows(archive, "/cropped/uncertain_lesion/crop_mapping.csv")
            for is_uncertain, mapping in [(False, main), (True, uncertain)]:
                for item in mapping:
                    original = old_by_id[Path(item["original_filename"]).stem]
                    filename = Path(item["cropped_filename"]).name
                    row = dict(image_id=filename.rsplit(".", 1)[0],
                               crop_path=f"cropped/prior_cohort/{filename}",
                               original_id=original["image_id"],
                               source_image_id=original["image_id"],
                               original_path=f"originals/prior_cohort/{original['image_id']}{Path(original['canonical_relative_path']).suffix.lower()}",
                               original_md5=original["md5"],
                               patient_group=original["patient_group"], label=int(original["label"]),
                               source_patient_group=original.get("source_patient_group", original["patient_group"]),
                               patient_identity_basis=("existing_manual_patient_link" if original["patient_group"].startswith("raw_patient_manual_")
                                                       else original.get("source_group_type", "unspecified")),
                               class_name=original["class_name"], cohort="prior_cohort",
                               source_batch=prior_batch(original), label_source="prior_clinical_manifest",
                               source_original_path=original["source_relative_path"],
                               source_package=package.name, crop_notes=item.get("notes", ""),
                               lesion_visibility=item.get("lesion_present", ""),
                               clinical_review="supplied_cropper_mapping_not_new_physician_confirmation")
                    reason = ""
                    if original["patient_group"] in overlap_groups:
                        reason = "patient_reserved_for_batch07"
                    elif is_uncertain and filename in main_names:
                        reason = "replaced_by_current_main_crop"
                    elif is_uncertain or item.get("lesion_present", "").lower() != "yes":
                        reason = "target_visibility_not_confirmed_in_supplied_mapping"
                    elif row["crop_path"] in quality_decisions:
                        decision = quality_decisions[row["crop_path"]]
                        reason = decision["reason"] if decision["decision"] == "exclude" else ""
                        row["framing_review"] = decision.get("evidence", "")
                    if reason:
                        excluded.append(dict(row, decision="exclude", reason=reason))
                        continue
                    suffix = "/cropped/" + filename
                    found = [n for n in members if n.endswith(suffix)]
                    if len(found) != 1:
                        raise ValueError(f"Ambiguous or missing crop: {package.name}:{filename}")
                    data = archive.read(found[0])
                    with Image.open(io.BytesIO(data)) as image:
                        image.load()
                        width, height = image.size
                    if (width, height) != (int(item["output_width"]), int(item["output_height"])):
                        raise ValueError(f"Crop dimension mismatch: {filename}")
                    row.update(crop_sha256=hashlib.sha256(data).hexdigest(),
                               crop_md5=hashlib.md5(data).hexdigest(), width=width, height=height,
                               crop_bytes=len(data), eligibility="included",
                               framing_review=row.get("framing_review", "contact_sheet_eye_framing_check"))
                    for field in ["crop_x", "crop_y", "crop_width", "crop_height"]:
                        row[field] = item.get(field, "")
                    records.append(row)
    return records, excluded, covered


def build_experiment_manifest(args):
    """Build lightweight metadata from supplied archives without extracting photos."""
    metadata = Path(args.metadata_root).resolve()
    old, patient_aliases = normalized_prior_patients(read_csv(args.old_manifest))
    old_by_id = {r["image_id"]: r for r in old}
    prior_exclusions = read_csv(metadata / "prior_development_exclusions.csv")
    overlap_groups = {r["patient_group"] for r in prior_exclusions}
    new_raw = read_csv(metadata / "batch07_originals.csv")
    new_by_md5 = {r["md5"]: r for r in new_raw}
    prior_md5_groups = defaultdict(set)
    prior_case_groups = defaultdict(set)
    for row in old:
        prior_md5_groups[row["md5"]].add(row["patient_group"])
        cid = case_id(row["source_relative_path"])
        if cid:
            prior_case_groups[cid].add(row["patient_group"])
    for original in new_raw:
        matched = prior_md5_groups[original["md5"]] | prior_case_groups[original["case_id"]]
        if len(matched) > 1:
            raise ValueError(f"Ambiguous patient identity: {original['case_id']}")
        for group in matched:
            labels = {int(r["label"]) for r in old if r["patient_group"] == group}
            if labels != {int(original["label"])}:
                raise ValueError(f"Prior and new clinical labels disagree: {original['case_id']}")
        overlap_groups.update(matched)
    decisions = read_csv(args.crop_decisions)
    by_path = {r["crop_path"]: r for r in decisions}
    if len(by_path) != len(decisions):
        raise ValueError("Repeated crop quality decision")
    packages = select_manual_packages(args.manual_packages)
    retrospective, exclusions, covered = package_crop_records(packages, old_by_id, overlap_groups, by_path)
    mapped_originals = {r["original_id"] for r in retrospective + exclusions}
    if covered != mapped_originals:
        raise ValueError("Every supplied manual-package original must have a returned crop mapping")
    # Author decision: retain these photos, but do not treat an anonymous
    # image group as an independently identified patient in formal splits.
    unidentified = [r for r in retrospective if r["patient_identity_basis"] == "unkeyed_singleton"]
    retrospective = [r for r in retrospective if r["patient_identity_basis"] != "unkeyed_singleton"]
    exclusions.extend(dict(r, decision="exclude", reason="patient_identity_unverified_preserve_image")
                      for r in unidentified)
    available_paths = {r["crop_path"] for r in retrospective} | {r["crop_path"] for r in exclusions}
    if set(by_path) - available_paths:
        raise ValueError("Crop decision does not match supplied packages")
    doctor_reviews = read_csv(metadata / "batch07_doctor_review.csv")
    temporal = []
    for crop in read_csv(metadata / "batch07_crops.csv"):
        original = new_by_md5[crop["original_md5"]]
        matched = prior_md5_groups[original["md5"]] | prior_case_groups[original["case_id"]]
        if len(matched) > 1:
            raise ValueError(f"Ambiguous patient identity: {original['case_id']}")
        group = next(iter(matched)) if matched else f"batch07_patient_{original['case_id']}"
        matching_reviews = [r for r in doctor_reviews
                            if r["original_filename"].replace("\\", "/") == crop["source_lookup_original"]]
        relative = crop["source_crop_path"].split("cropped/batch07/", 1)[-1]
        temporal.append(dict(image_id=f"crop_{crop['crop_sha256'][:20]}",
                             crop_path=f"cropped/batch07/{relative}",
                             original_id=f"img_{original['sha256'][:20]}",
                             source_image_id=original.get("prior_image_ids", ""),
                             original_path=f"originals/batch07/{original['source_lookup_path']}",
                             original_md5=original["md5"], patient_group=group,
                             patient_identity_basis="supplied_clinical_case_directory",
                             label=int(crop["label"]), class_name=crop["class_name"],
                             cohort="batch07", source_batch=7,
                             label_source=crop["label_source"],
                             source_original_path=original["source_lookup_path"],
                             source_package="batch07_supplied_folders", crop_notes=crop["mapping_notes"],
                             lesion_visibility="confirmed_by_supplied_manual_mapping",
                             clinical_review="explicit_supplied_doctor_review" if matching_reviews else "supplied_manual_mapping",
                             doctor_comment=";".join(r.get("审核意见", "") for r in matching_reviews),
                             crop_sha256=crop["crop_sha256"], crop_md5=crop["crop_md5"],
                             width=int(crop["width"]), height=int(crop["height"]), crop_bytes=int(crop["bytes"]),
                             eligibility="included", split="prospective_test",
                             framing_review="contact_sheet_eye_framing_check"))
        for field in ["crop_x", "crop_y", "crop_width", "crop_height"]:
            temporal[-1][field] = crop.get(field, "")
    missing = [r for r in old if r["image_id"] not in covered and r["patient_group"] not in overlap_groups]
    crop_policy = args.crop_source_policy
    if crop_policy == "existing_manual_only" and args.recovery_plan:
        raise ValueError("A recovery plan cannot be activated under the existing-manual-only policy")
    missing_plan = read_csv(args.recovery_plan) if args.recovery_plan else []
    by_original = {r["image_id"]: r for r in missing_plan}
    if by_original and set(by_original) != {r["image_id"] for r in missing}:
        raise ValueError("Recovery plan does not cover exactly the missing originals")
    missing_crop_actions = []
    for original in missing:
        plan = by_original.get(original["image_id"], {})
        action = plan.get("action", "recover_eye_crop_on_remote")
        reason = plan.get("reason", "original_not_in_returned_manual_crop_packages")
        if crop_policy == "existing_manual_only":
            action = "preserve_original_only_no_new_crop"
            reason = "not_in_returned_manual_packages_no_additional_cropping_requested"
        missing_crop_actions.append(dict(original_id=original["image_id"],
                                         original_md5=original["md5"], patient_group=original["patient_group"],
                                         label=int(original["label"]), label_source="prior_clinical_manifest",
                                         original_path=f"originals/prior_cohort/{original['image_id']}{Path(original['canonical_relative_path']).suffix.lower()}",
                                         source_original_path=original["source_relative_path"],
                                         action=action, reason=reason))
    allocation = grouped_partition(retrospective, args.split_seed)
    for row in retrospective:
        row["split"] = allocation[row["patient_group"]]
    included = retrospective + temporal
    crossings = assert_partition_integrity(included)
    patient_rows = []
    for group in sorted({r["patient_group"] for r in included}):
        items = [r for r in included if r["patient_group"] == group]
        patient_rows.append(dict(patient_group=group, label=items[0]["label"], split=items[0]["split"],
                                 cohort=items[0]["cohort"], crop_count=len(items),
                                 original_count=len({r["original_md5"] for r in items})))
    counts = []
    for split in (*SPLIT_NAMES, "prospective_test"):
        items = [r for r in included if r["split"] == split]
        patients = [r for r in patient_rows if r["split"] == split]
        counts.append(dict(split=split, crops=len(items), patients=len(patients),
                           originals=len({r["original_md5"] for r in items}),
                           benign_crops=sum(int(r["label"]) == 0 for r in items),
                           malignant_crops=sum(int(r["label"]) == 1 for r in items),
                           benign_patients=sum(int(r["label"]) == 0 for r in patients),
                           malignant_patients=sum(int(r["label"]) == 1 for r in patients)))
    remaining = (sum(r["action"] != "exclude" for r in missing_crop_actions)
                 if crop_policy == "supplement_missing_originals" else 0)
    summary = dict(retrospective_batches=[1, 2, 3, 4, 5, 6], retrospective_patient_ratios=[0.8, 0.1, 0.1],
                   split_seed=args.split_seed, assignment_unit="patient_group", stratification="clinical_class",
                   later_cohort_batch=7, later_cohort_requested_name="Prospective Cohort Test",
                   later_cohort_actual_design="same_hospital_later_batch_temporal_validation",
                   collection_dates_supplied=False, prospective_enrollment_evidence_supplied=False,
                   prospective_status_verified=False,
                   included_crops=len(included), retrospective_crops=len(retrospective),
                   later_cohort_crops=len(temporal), counts=counts,
                   identity_crossings=crossings, excluded_crop_records=len(exclusions),
                   prior_patient_groups_reserved_for_batch07=len(overlap_groups),
                   exclusion_reasons=dict(Counter(r["reason"] for r in exclusions)),
                   crop_source_policy=crop_policy, new_crop_creation_enabled=crop_policy == "supplement_missing_originals",
                   manual_package_originals=len(covered),
                   manual_package_crop_records=len(retrospective) + sum(r["cohort"] == "prior_cohort" for r in exclusions),
                   manual_package_mapping_complete=covered == mapped_originals,
                   source_inventory_coverage_complete=not missing,
                   source_archives=[dict(path=str(p.resolve()), bytes=p.stat().st_size) for p in packages],
                   originals_without_returned_manual_crop=len(missing),
                   originals_needing_remote_crop_creation=remaining,
                   supplied_doctor_excluded_batch07_originals=len(read_csv(metadata / "batch07_originals_without_crop.csv")),
                   remote_pixel_validation_complete=False,
                   dataset_complete=False, training_ready=False,
                   state=("existing_manual_crops_selected_remote_file_validation_required"
                          if crop_policy == "existing_manual_only" else
                          "metadata_prepared_remote_crop_recovery_and_file_validation_required"))
    write_csv(metadata / "dataset_manifest.csv", included)
    write_csv(metadata / "patient_split.csv", patient_rows)
    write_csv(metadata / "split_counts.csv", counts)
    write_csv(metadata / "crop_exclusions.csv", exclusions)
    write_csv(metadata / "missing_crop_actions.csv", missing_crop_actions)
    write_json(metadata / "cohort_design.json", summary)
    write_csv(metadata / "patient_group_aliases.csv",
              [dict(source_patient_group=group, patient_group=target,
                    reason="same_clinical_case_in_dated_followup_subfolder")
               for group, target in sorted(patient_aliases.items())],
              ["source_patient_group", "patient_group", "reason"])
    print(json.dumps(summary, ensure_ascii=False))
    return 0


def prepare_local_split_delivery(args):
    """Keep one copy of each selected manual crop and export four usable splits."""
    workspace = Path(args.workspace).resolve()
    dataset = Path(args.dataset_root).resolve()
    if dataset != workspace / "eye_dataset":
        raise ValueError("Local dataset must use the existing project eye_dataset directory")
    metadata = dataset / "metadata"
    rows = read_csv(metadata / "dataset_manifest.csv")
    design = json.loads((metadata / "cohort_design.json").read_text(encoding="utf-8"))
    if design.get("crop_source_policy") != "existing_manual_only":
        raise ValueError("Local delivery requires the existing-manual-only crop policy")
    crossings = assert_partition_integrity(rows)
    archives = {p.name: p for p in select_manual_packages(args.manual_packages)}
    output = []
    counts = Counter()
    verified_originals = {}
    with ExitStack() as stack:
        opened = {name: stack.enter_context(zipfile.ZipFile(path)) for name, path in archives.items()}
        members = {name: [n for n in z.namelist() if not hidden_zip_entry(n)] for name, z in opened.items()}
        for row in rows:
            destination = safe_relative(dataset, row["crop_path"])
            filename = destination.name
            original_archive, original_member, original_local = "", "", ""
            if row["cohort"] == "prior_cohort":
                name = row["source_package"]
                suffix = "/cropped/" + filename
                matches = [n for n in members[name] if n.endswith(suffix)]
                if len(matches) != 1:
                    raise ValueError(f"Missing or ambiguous manual crop: {filename}")
                data = opened[name].read(matches[0])
                if hashlib.sha256(data).hexdigest() != row["crop_sha256"]:
                    raise ValueError(f"Manual crop changed after manifest creation: {filename}")
                original_matches = [n for n in members[name] if "/originals/" in n
                                    and Path(n).stem == row["source_image_id"]
                                    and Path(n).suffix.lower() in IMAGE_EXTS]
                if len(original_matches) != 1:
                    raise ValueError(f"Missing or ambiguous original in archive: {filename}")
                original_archive, original_member = str(archives[name]), original_matches[0]
                key = (original_archive, original_member)
                if key not in verified_originals:
                    verified_originals[key] = hashlib.md5(opened[name].read(original_member)).hexdigest()
                if verified_originals[key] != row["original_md5"]:
                    raise ValueError(f"Archive original does not match the clinical source: {filename}")
                destination.parent.mkdir(parents=True, exist_ok=True)
                if destination.exists():
                    if hashes(destination)[1] != row["crop_sha256"]:
                        raise ValueError(f"Existing crop differs; refusing overwrite: {destination}")
                    counts["existing_prior_crops"] += 1
                else:
                    with destination.open("xb") as stream:
                        stream.write(data)
                    counts["extracted_manual_crops"] += 1
            else:
                relative = row["crop_path"].split("cropped/batch07/", 1)[1]
                source = safe_relative(workspace, relative)
                if not source.is_file():
                    source = safe_relative(dataset, row["crop_path"])
                if hashes(source)[1] != row["crop_sha256"]:
                    raise ValueError(f"Batch-seven manual crop changed: {filename}")
                if not destination.exists():
                    link_file(source, destination)
                    counts["batch07_hardlinks"] += 1
                elif hashes(destination)[1] != row["crop_sha256"]:
                    raise ValueError(f"Existing crop differs; refusing overwrite: {destination}")
                else:
                    counts["existing_batch07_crops"] += 1
                original_relative = row["original_path"].split("originals/batch07/", 1)[1]
                original = safe_relative(workspace, original_relative)
                if not original.is_file():
                    original = safe_relative(dataset, row["original_path"])
                if not original.is_file():
                    raise FileNotFoundError(original)
                original_local = str(original)
                key = (original_local, "")
                if key not in verified_originals:
                    verified_originals[key] = hashes(original)[0]
                if verified_originals[key] != row["original_md5"]:
                    raise ValueError(f"Batch-seven original does not match its source: {filename}")
            info = image_info(destination)
            if (info["width"], info["height"]) != (int(row["width"]), int(row["height"])):
                raise ValueError(f"Native crop dimensions changed: {filename}")
            if hashes(destination)[1] != row["crop_sha256"]:
                raise ValueError(f"Stored crop hash differs: {filename}")
            with Image.open(destination) as image:
                tiny = list(image.convert("L").resize((9, 8), Image.Resampling.LANCZOS).getdata())
            bits = [tiny[y * 9 + x + 1] > tiny[y * 9 + x] for y in range(8) for x in range(8)]
            crop_dhash = f"{int(''.join('1' if bit else '0' for bit in bits), 2):016x}"
            output.append(dict(image_id=row["image_id"], patient_group=row["patient_group"],
                               cohort="Prospective Cohort" if row["cohort"] == "batch07" else "Retrospective Cohort",
                               split=row["split"], source_batch=row["source_batch"],
                               label=int(row["label"]), class_name=row["class_name"],
                               crop_path=row["crop_path"], original_id=row["original_id"],
                               original_md5=row["original_md5"], crop_sha256=row["crop_sha256"],
                               crop_dhash=crop_dhash,
                               width=int(row["width"]), height=int(row["height"]),
                               patient_identity_basis=row["patient_identity_basis"],
                               label_source=row["label_source"], original_archive=original_archive,
                               original_archive_member=original_member, original_local_path=original_local))
            if len(output) % 500 == 0:
                print(json.dumps({"validated_crops": len(output), "total_crops": len(rows)}), flush=True)
    names = {"train": "retrospective_train.csv", "validation": "retrospective_validation.csv",
             "test": "retrospective_test.csv", "prospective_test": "prospective_test.csv"}
    split_root = dataset / "splits"
    split_root.mkdir(parents=True, exist_ok=True)
    exported = {}
    for split, filename in names.items():
        part = sorted((r for r in output if r["split"] == split), key=lambda r: r["image_id"])
        path = split_root / filename
        write_csv(path, part)
        exported[filename] = dict(crops=len(part), patient_groups=len({r["patient_group"] for r in part}),
                                  sha256=hashes(path)[1])
    write_csv(split_root / "split_summary.csv", design["counts"])
    unkeyed = [r for r in output if r["patient_identity_basis"] == "unkeyed_singleton"]
    summary = dict(author="Jia Rong", split_assignment_complete=True, local_crop_files_verified=True,
                   local_dataset_root=str(dataset), selected_manual_crops=len(output),
                   verified_original_files=len(verified_originals), original_source_hashes_match=True,
                   split_seed=design["split_seed"], retrospective_patient_ratios=[0.8, 0.1, 0.1],
                   label_encoding={"benign": 0, "malignant": 1},
                   image_operations="original_manual_crop_bytes_preserved",
                   originals_storage="prior_originals_in_existing_archives_batch07_in_supplied_folders",
                   identity_crossings=crossings, files=exported, counts=design["counts"],
                   prior_patient_groups_reserved_for_batch07=design["prior_patient_groups_reserved_for_batch07"],
                   unresolved_patient_identity_groups=len({r["patient_group"] for r in unkeyed}),
                   crops_with_unresolved_patient_identity=len(unkeyed),
                   patient_level_independence_fully_verified=not unkeyed,
                   later_cohort_display_name="Prospective Cohort Test",
                   later_cohort_design="same_hospital_later_batch_temporal_validation",
                   prospective_enrollment_verified=False,
                   remote_sync_complete=False, file_operations=dict(counts))
    write_json(split_root / "split_definition.json", summary)
    design.update(split_assignment_complete=True, local_pixel_validation_complete=True,
                  selected_manual_dataset_complete=True,
                  unresolved_patient_identity_groups=summary["unresolved_patient_identity_groups"],
                  crops_with_unresolved_patient_identity=len(unkeyed),
                  patient_level_independence_fully_verified=not unkeyed,
                  state="local_manual_crop_splits_prepared",
                  training_ready=not unkeyed)
    write_json(metadata / "cohort_design.json", design)
    print(json.dumps(summary, ensure_ascii=False))
    return 0


def verify_local_split_delivery(args):
    """Independently reconcile exported indexes, master metadata, and crop bytes."""
    dataset = Path(args.dataset_root).resolve()
    split_root = dataset / "splits"
    definition = json.loads((split_root / "split_definition.json").read_text(encoding="utf-8"))
    master = {r["image_id"]: r for r in read_csv(dataset / "metadata" / "dataset_manifest.csv")}
    filenames = {"train": "retrospective_train.csv", "validation": "retrospective_validation.csv",
                 "test": "retrospective_test.csv", "prospective_test": "prospective_test.csv"}
    observed = {}
    rows = []
    for split, filename in filenames.items():
        index = split_root / filename
        if hashes(index)[1] != definition["files"][filename]["sha256"]:
            raise ValueError(f"Exported index has changed: {filename}")
        current = read_csv(index)
        for row in current:
            identity = row["image_id"]
            if identity in observed or identity not in master:
                raise ValueError(f"Repeated or unknown exported image ID: {identity}")
            expected = master[identity]
            for field in ["split", "patient_group", "label", "crop_path", "original_md5", "crop_sha256", "source_batch"]:
                if row[field] != expected[field]:
                    raise ValueError(f"Export disagrees with master metadata: {identity}:{field}")
            if row["split"] != split:
                raise ValueError(f"Record is in the wrong split file: {identity}")
            path = safe_relative(dataset, row["crop_path"])
            if hashes(path)[1] != row["crop_sha256"]:
                raise ValueError(f"Image differs from exported index: {identity}")
            with Image.open(path) as image:
                image.load()
                if image.size != (int(row["width"]), int(row["height"])):
                    raise ValueError(f"Unexpected image dimensions: {identity}")
            observed[identity] = row
            rows.append(dict(row, cohort="batch07" if split == "prospective_test" else "prior_cohort"))
        if len(current) != definition["files"][filename]["crops"]:
            raise ValueError(f"Split count differs: {filename}")
    if set(observed) != set(master):
        raise ValueError("Exported splits do not cover exactly the selected master dataset")
    actual_paths = {p.relative_to(dataset).as_posix() for p in (dataset / "cropped").rglob("*")
                    if p.is_file() and p.suffix.lower() in IMAGE_EXTS}
    indexed_paths = {r["crop_path"] for r in rows}
    retained = {r["crop_path"]: r for r in read_csv(dataset / "metadata" / "crop_exclusions.csv")
                if r.get("reason") == "patient_identity_unverified_preserve_image"}
    extras = actual_paths - indexed_paths
    if indexed_paths - actual_paths or extras - set(retained):
        raise ValueError("Stored crop directory and index differ")
    for relative in extras:
        if hashes(safe_relative(dataset, relative))[1] != retained[relative]["crop_sha256"]:
            raise ValueError(f"Preserved excluded crop changed: {relative}")
    crossings = assert_partition_integrity(rows)
    definition["independent_index_check"] = dict(passed=True, crops=len(rows),
                                               all_files_decoded=True, all_hashes_match=True,
                                               no_duplicate_ids=True, no_unaccounted_crop_files=True,
                                               preserved_excluded_crops=len(extras),
                                               identity_crossings=crossings)
    write_json(split_root / "split_definition.json", definition)
    print(json.dumps(definition["independent_index_check"]))
    return 0


def pixel_box(relative_box, size):
    if len(relative_box) != 4:
        raise ValueError("Each crop box requires four coordinates")
    left, top, right, bottom = map(float, relative_box)
    if not (0 <= left < right <= 1 and 0 <= top < bottom <= 1):
        raise ValueError(f"Invalid normalized crop box: {relative_box}")
    width, height = size
    box = (round(left * width), round(top * height),
           round(right * width), round(bottom * height))
    if box[0] >= box[2] or box[1] >= box[3]:
        raise ValueError("Crop is empty at the native image resolution")
    return box


def rotate_without_resampling(image, clockwise):
    methods = {0: None, 90: Image.Transpose.ROTATE_270,
               180: Image.Transpose.ROTATE_180, 270: Image.Transpose.ROTATE_90}
    if clockwise not in methods:
        raise ValueError("Rotation must be a multiple of 90 degrees")
    return image if clockwise == 0 else image.transpose(methods[clockwise])


def recover_missing_crops(args):
    """Run on the canonical remote data, preserving all original images."""
    project = Path(args.project_root).resolve()
    dataset = Path(args.dataset_root).resolve()
    if not dataset.is_relative_to(project / "data_stream"):
        raise ValueError("Only this project's canonical data_stream dataset is allowed")
    check_crop_creation_policy(dataset)
    plan = json.loads(Path(args.recovery_plan).read_text(encoding="utf-8"))
    rows = []
    for instruction in plan["records"]:
        source = safe_relative(dataset, instruction["original_path"])
        md5, original_sha = hashes(source)
        if md5 != instruction["original_md5"]:
            raise ValueError(f"Original MD5 differs from the reviewed preview source: {source}")
        if instruction["action"] == "exclude":
            continue
        with Image.open(source) as original:
            original.load()
            image = ImageOps.exif_transpose(original).convert("RGB")
            for index, proposal in enumerate(instruction["boxes"], 1):
                box = pixel_box(proposal, image.size)
                rotation = int(instruction["clockwise_rotation"])
                filename = f"{instruction['image_id']}_recovered_eye_{index:02d}.jpg"
                destination = dataset / "cropped" / "prior_cohort" / filename
                destination.parent.mkdir(parents=True, exist_ok=True)
                output = rotate_without_resampling(image.crop(box), rotation)
                if instruction["action"] == "use_native_photo_panel":
                    if box != (0, 0, image.width, image.height) or rotation != 0:
                        raise ValueError("Native photo panel must not be altered")
                    link_file(source, destination)
                else:
                    buffer = io.BytesIO()
                    output.save(buffer, format="JPEG", quality=95, subsampling=0)
                    data = buffer.getvalue()
                    if destination.exists():
                        if hashlib.sha256(data).hexdigest() != hashes(destination)[1]:
                            raise ValueError(f"Existing recovered crop differs; no overwrite: {destination}")
                    else:
                        destination.write_bytes(data)
                crop_md5, crop_sha = hashes(destination)
                record = dict(image_id=Path(filename).stem,
                              crop_path=destination.relative_to(dataset).as_posix(),
                              original_id=f"img_{original_sha[:20]}",
                              source_image_id=instruction["image_id"],
                              original_path=instruction["original_path"], original_md5=md5,
                              patient_group=instruction["patient_group"], label=instruction["label"],
                              class_name="malignant" if int(instruction["label"]) else "benign",
                              cohort="prior_cohort", label_source=instruction["label_source"],
                              source_original_path=instruction["source_original_path"],
                              crop_sha256=crop_sha, crop_md5=crop_md5, crop_bytes=destination.stat().st_size,
                              crop_x=box[0], crop_y=box[1], crop_width=box[2]-box[0], crop_height=box[3]-box[1],
                              clockwise_rotation=rotation, crop_type="eye",
                              recovery_reason=instruction["reason"],
                              eligibility="requires_native_crop_framing_check",
                              clinical_review="clinical_class_inherited_not_new_physician_confirmation",
                              **image_info(destination))
                rows.append(record)
        print(json.dumps({"recovered_original": instruction["image_id"], "crop_count": len(rows)}), flush=True)
    write_csv(dataset / "metadata" / "recovered_crops.csv", rows)
    write_json(dataset / "metadata" / "crop_recovery_status.json",
               dict(original_instructions=len(plan["records"]), produced_crops=len(rows),
                    excluded_originals=sum(r["action"] == "exclude" for r in plan["records"]),
                    native_visual_review_complete=False, training_ready=False,
                    no_upscaling=True, intensity_enhancement=False))
    print(json.dumps({"produced_crops": len(rows), "native_visual_review_complete": False}))
    return 0


def validate_remote_manifest(args):
    dataset = Path(args.dataset_root).resolve()
    metadata = dataset / "metadata"
    rows = read_csv(metadata / "dataset_manifest.csv")
    crossings = assert_partition_integrity(rows)
    originals = read_csv(metadata / "originals.csv")
    original_catalog = {r["md5"]: r for r in originals}
    if len(original_catalog) != len(originals):
        raise ValueError("Canonical original catalog contains repeated MD5 identities")
    original_hashes = {}
    for row in rows:
        canonical = original_catalog.get(row["original_md5"])
        if canonical is None:
            raise ValueError(f"Original absent from the canonical dataset catalog: {row['image_id']}")
        if (canonical["patient_group"], int(canonical["label"])) != (row["patient_group"], int(row["label"])):
            raise ValueError(f"Canonical patient/clinical label differs: {row['image_id']}")
        row["original_id"] = canonical["original_id"]
        row["original_path"] = canonical["original_path"]
        crop = safe_relative(dataset, row["crop_path"])
        original = safe_relative(dataset, row["original_path"])
        if hashes(crop)[1] != row["crop_sha256"]:
            raise ValueError(f"Current crop differs from the supplied manifest: {crop}")
        if original not in original_hashes:
            original_hashes[original] = hashes(original)[0]
        if original_hashes[original] != row["original_md5"]:
            raise ValueError(f"Current original differs from the supplied manifest: {original}")
        info = image_info(crop)
        if (info["width"], info["height"]) != (int(row["width"]), int(row["height"])):
            raise ValueError(f"Crop dimension differs from the manifest: {crop}")
    summary = dict(current_manifest_crops=len(rows), verified_original_files=len(original_hashes),
                   identity_crossings=crossings, original_and_crop_hashes_match=True,
                   all_included_crops_decoded=True, original_ids_match_canonical_catalog=True)
    write_csv(metadata / "dataset_manifest.csv", rows)
    write_json(metadata / "current_manifest_file_validation.json", summary)
    design = json.loads((metadata / "cohort_design.json").read_text(encoding="utf-8"))
    design["remote_pixel_validation_complete"] = True
    design["training_ready"] = bool(design["dataset_complete"])
    write_json(metadata / "cohort_design.json", design)
    print(json.dumps(summary))
    return 0


def finalize_recovered_manifest(args):
    """Include recovered crops only after a complete explicit framing review."""
    metadata = Path(args.metadata_root).resolve()
    rows = read_csv(metadata / "dataset_manifest.csv")
    recovered = read_csv(metadata / "recovered_crops.csv")
    review = read_csv(args.framing_review)
    decisions = {r["crop_path"]: r for r in review}
    if len(decisions) != len(review):
        raise ValueError("Repeated recovered-crop framing decision")
    if set(decisions) != {r["crop_path"] for r in recovered}:
        raise ValueError("Every recovered crop requires an explicit framing decision")
    normalized, _ = normalized_prior_patients(read_csv(args.old_manifest))
    old = {r["image_id"]: r for r in normalized}
    exclusions = read_csv(metadata / "crop_exclusions.csv")
    included_paths = {r["crop_path"] for r in rows}
    for row in recovered:
        decision = decisions[row["crop_path"]]
        if decision["decision"] not in ("include", "exclude"):
            raise ValueError("Framing decisions must be include or exclude, never pending")
        source_id = row.get("source_image_id", row["original_id"])
        original = old[source_id]
        if (row["patient_group"], int(row["label"]), row["original_md5"]) != (
                original["patient_group"], int(original["label"]), original["md5"]):
            raise ValueError(f"Recovered crop clinical/source identity mismatch: {row['crop_path']}")
        row.update(source_batch=prior_batch(original),
                   framing_review=decision.get("evidence", "explicit_native_crop_framing_review"),
                   lesion_visibility=decision.get("target_visibility", "visible_in_reviewed_crop"),
                   source_package="canonical_original_cpu_crop_recovery")
        if decision["decision"] == "exclude":
            exclusions.append(dict(row, decision="exclude", reason=decision["reason"]))
        elif row["crop_path"] not in included_paths:
            row["eligibility"] = "included"
            rows.append(row)
            included_paths.add(row["crop_path"])
    retrospective = [r for r in rows if r["cohort"] == "prior_cohort"]
    temporal = [r for r in rows if r["cohort"] == "batch07"]
    allocation = grouped_partition(retrospective, args.split_seed)
    for row in retrospective:
        row["split"] = allocation[row["patient_group"]]
    for row in temporal:
        row["split"] = "prospective_test"
    crossings = assert_partition_integrity(rows)
    patient_rows, counts = [], []
    patient_items = defaultdict(list)
    for row in rows:
        patient_items[row["patient_group"]].append(row)
    for group, items in sorted(patient_items.items()):
        patient_rows.append(dict(patient_group=group, label=items[0]["label"], split=items[0]["split"],
                                 cohort=items[0]["cohort"], crop_count=len(items),
                                 original_count=len({r["original_md5"] for r in items})))
    for split in (*SPLIT_NAMES, "prospective_test"):
        items = [r for r in rows if r["split"] == split]
        patients = [r for r in patient_rows if r["split"] == split]
        counts.append(dict(split=split, crops=len(items), patients=len(patients),
                           originals=len({r["original_md5"] for r in items}),
                           benign_crops=sum(int(r["label"]) == 0 for r in items),
                           malignant_crops=sum(int(r["label"]) == 1 for r in items),
                           benign_patients=sum(int(r["label"]) == 0 for r in patients),
                           malignant_patients=sum(int(r["label"]) == 1 for r in patients)))
    actions = read_csv(metadata / "missing_crop_actions.csv")
    recovered_originals = {r.get("source_image_id", r["original_id"]) for r in recovered}
    if any(r["action"] != "exclude" and r["original_id"] not in recovered_originals for r in actions):
        raise ValueError("Some originals still require crop recovery")
    design = json.loads((metadata / "cohort_design.json").read_text(encoding="utf-8"))
    design.update(counts=counts, included_crops=len(rows), retrospective_crops=len(retrospective),
                  later_cohort_crops=len(temporal), identity_crossings=crossings,
                  excluded_crop_records=len(exclusions),
                  originals_needing_remote_crop_creation=0,
                  dataset_complete=True, remote_pixel_validation_complete=False,
                  training_ready=False, state="crop_review_complete_file_validation_required")
    write_csv(metadata / "dataset_manifest.csv", rows)
    write_csv(metadata / "patient_split.csv", patient_rows)
    write_csv(metadata / "split_counts.csv", counts)
    write_csv(metadata / "crop_exclusions.csv", exclusions)
    write_json(metadata / "cohort_design.json", design)
    print(json.dumps(design, ensure_ascii=False))
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    local = commands.add_parser("audit-new")
    local.add_argument("--workspace", required=True)
    local.add_argument("--old-manifest", required=True)
    local.add_argument("--output", required=True)
    local.add_argument("--preview-dir")
    local.set_defaults(func=audit_new)
    remote = commands.add_parser("consolidate-remote")
    remote.add_argument("--project-root", required=True)
    remote.add_argument("--dataset-root", required=True)
    remote.add_argument("--old-manifest", required=True)
    remote.add_argument("--manual-packages", required=True)
    remote.set_defaults(func=consolidate_remote)
    preview = commands.add_parser("preview-catalog")
    preview.add_argument("--dataset-root", required=True)
    preview.add_argument("--family", choices=["main", "uncertain", "missing"], required=True)
    preview.add_argument("--output", required=True)
    preview.set_defaults(func=preview_catalog)
    manifest = commands.add_parser("build-experiment-manifest")
    manifest.add_argument("--metadata-root", required=True)
    manifest.add_argument("--old-manifest", required=True)
    manifest.add_argument("--manual-packages", required=True)
    manifest.add_argument("--crop-decisions", required=True)
    manifest.add_argument("--recovery-plan")
    manifest.add_argument("--crop-source-policy", choices=["existing_manual_only", "supplement_missing_originals"],
                          default="existing_manual_only")
    manifest.add_argument("--split-seed", type=int, default=42)
    manifest.set_defaults(func=build_experiment_manifest)
    delivery = commands.add_parser("prepare-local-splits")
    delivery.add_argument("--workspace", required=True)
    delivery.add_argument("--dataset-root", required=True)
    delivery.add_argument("--manual-packages", required=True)
    delivery.set_defaults(func=prepare_local_split_delivery)
    local_check = commands.add_parser("verify-local-splits")
    local_check.add_argument("--dataset-root", required=True)
    local_check.set_defaults(func=verify_local_split_delivery)
    recovery = commands.add_parser("recover-missing-crops")
    recovery.add_argument("--project-root", required=True)
    recovery.add_argument("--dataset-root", required=True)
    recovery.add_argument("--recovery-plan", required=True)
    recovery.set_defaults(func=recover_missing_crops)
    check = commands.add_parser("validate-current-manifest")
    check.add_argument("--dataset-root", required=True)
    check.set_defaults(func=validate_remote_manifest)
    final = commands.add_parser("finalize-recovered-manifest")
    final.add_argument("--metadata-root", required=True)
    final.add_argument("--old-manifest", required=True)
    final.add_argument("--framing-review", required=True)
    final.add_argument("--split-seed", type=int, default=42)
    final.set_defaults(func=finalize_recovered_manifest)
    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
